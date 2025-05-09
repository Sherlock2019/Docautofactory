from docx import Document
import os

def replace_placeholders_and_generate(template_stream, replacements, output_path):
    doc = Document(template_stream)

    for para in doc.paragraphs:
        for key, val in replacements.items():
            if key in para.text:
                if val and isinstance(val, str):
                    new_text = para.text.replace(f"{{{key}}}", f"<<[{os.path.basename(val)}]>>" if os.path.isfile(val) else val)
                    para.text = new_text

    doc.save(output_path)

